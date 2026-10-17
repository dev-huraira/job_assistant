from pathlib import Path

from docx import Document

from app.core.logging_config import setup_logging
from app.services.cv_parser import CVParsingError, extract_cv_text

setup_logging()

# --- Build a sample .docx on the fly ---
docx_path = Path("sample_cv.docx")
document = Document()
document.add_paragraph("Jane Doe")
document.add_paragraph("Software Engineer")
document.add_paragraph("Experience: Built REST APIs with Python and Flask.")
document.save(docx_path)

docx_bytes = docx_path.read_bytes()
text = extract_cv_text("sample_cv.docx", docx_bytes)
print("=== DOCX extraction ===")
print(text)
print()

# --- Build a sample .txt file ---
txt_path = Path("sample_cv.txt")
txt_path.write_text("Jane Doe\nSoftware Engineer\nSkills: Python, Flask", encoding="utf-8")

txt_bytes = txt_path.read_bytes()
text = extract_cv_text("sample_cv.txt", txt_bytes)
print("=== TXT extraction ===")
print(text)
print()

# --- Try a real PDF, if you place one named sample_cv.pdf in this folder ---
pdf_path = Path("sample_cv.pdf")
if pdf_path.exists():
    pdf_bytes = pdf_path.read_bytes()
    text = extract_cv_text("sample_cv.pdf", pdf_bytes)
    print("=== PDF extraction (first 500 chars) ===")
    print(text[:500], "...")
else:
    print(
        "No sample_cv.pdf found -- copy a real PDF resume into this folder "
        "as 'sample_cv.pdf' to test PDF extraction."
    )
print()

# --- Error handling: unsupported extension ---
try:
    extract_cv_text("resume.exe", b"not a real cv")
    print("ERROR: should have raised CVParsingError!")
except CVParsingError as e:
    print("Correctly rejected unsupported file type:", e)

# --- Error handling: file too large ---
try:
    huge_bytes = b"x" * (6 * 1024 * 1024)
    extract_cv_text("huge.txt", huge_bytes)
    print("ERROR: should have raised CVParsingError for size!")
except CVParsingError as e:
    print("Correctly rejected oversized file:", e)